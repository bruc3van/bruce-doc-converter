import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from subprocess import CompletedProcess
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from bruce_doc_converter import cli
from bruce_doc_converter.converter import convert_document, setup_node_dependencies


class ReviewRegressions(unittest.TestCase):
    def test_setup_accepts_minimum_cli_runtime_with_ready_dependencies(self):
        with patch('bruce_doc_converter.converter._shared_node_dependencies_ready', return_value=True), \
             patch('bruce_doc_converter.converter.shutil.which', return_value='node'), \
             patch('bruce_doc_converter.converter.subprocess.run', return_value=CompletedProcess([], 0, 'v22.0.0', '')), \
             patch('bruce_doc_converter.converter._ensure_shared_node_modules') as install:
            result = setup_node_dependencies()
        self.assertTrue(result['success'], result)
        self.assertTrue(result['already_installed'])
        install.assert_not_called()

    def test_setup_rejects_missing_and_old_node_even_with_ready_dependencies(self):
        for ready in (False, True):
            for node, version, expected in [(None, '', 'NODE_NOT_FOUND'),
                                             ('node', 'v20.20.0', 'NODE_VERSION_UNSUPPORTED')]:
                with self.subTest(ready=ready, node=node), \
                     patch('bruce_doc_converter.converter._shared_node_dependencies_ready', return_value=ready), \
                     patch('bruce_doc_converter.converter.shutil.which', return_value=node), \
                     patch('bruce_doc_converter.converter.subprocess.run', return_value=CompletedProcess([], 0, version, '')), \
                     patch('bruce_doc_converter.converter._ensure_shared_node_modules') as install:
                    result = setup_node_dependencies()
                    self.assertFalse(result['success'], result)
                    self.assertEqual(expected, result['error_code'])
                    install.assert_not_called()

    def test_setup_cli_emits_runtime_recovery_suggestion(self):
        with patch('bruce_doc_converter.cli.setup_node_dependencies', return_value={
            'success': False, 'error_code': 'NODE_VERSION_UNSUPPORTED', 'error': 'old node'
        }), contextlib.redirect_stdout(io.StringIO()) as stdout:
            code = cli.main(['setup-node'])
        self.assertEqual(1, code)
        self.assertIn('22.0', json.loads(stdout.getvalue())['suggestion'])

    def test_batch_missing_directory_code_matches_with_and_without_auto_manifest(self):
        with tempfile.TemporaryDirectory() as td:
            missing = str(Path(td) / 'missing')
            codes = []
            for options in ([], ['--manifest', 'auto']):
                with contextlib.redirect_stdout(io.StringIO()) as stdout:
                    self.assertEqual(1, cli.main(['batch', missing, *options]))
                result = json.loads(stdout.getvalue())
                codes.append(result['error_code'] if 'error_code' in result else result['results'][0]['result']['error_code'])
            self.assertEqual(['FILE_NOT_FOUND', 'FILE_NOT_FOUND'], codes)

    def test_office_image_save_failure_warns_and_strict_rejects(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            png = root / 'picture.png'
            Image.new('RGB', (200, 150), 'red').save(png, compress_level=0)
            doc = Document(); doc.add_paragraph('body'); doc.add_picture(str(png)); doc.save(root / 'input.docx')
            book = Workbook(); book.active['A1'] = 'body'; book.active.add_image(SheetImage(str(png)), 'B2'); book.save(root / 'input.xlsx')
            slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[6])
            slide.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(2)); slides.save(root / 'input.pptx')
            for extension in ('docx', 'xlsx', 'pptx'):
                for strict in (False, True):
                    with self.subTest(extension=extension, strict=strict):
                        out = root / f'{extension}-{strict}'
                        with patch(f'bruce_doc_converter.formats.{extension}._save_extracted_image', return_value=None):
                            result = convert_document(root / f'input.{extension}', output_dir=out, strict=strict)
                        self.assertEqual(not strict, result['success'], result)
                        self.assertTrue(result['warnings'])
                        self.assertEqual('IMAGE_EXTRACTION_FAILED', result['diagnostics'][-1]['code'])
                        if strict:
                            self.assertEqual('CONTENT_INCOMPLETE', result['error_code'])
                            self.assertFalse(list(out.glob('*.md')))
                            self.assertFalse(list((out / 'images').iterdir()))
