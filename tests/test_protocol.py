import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from bruce_doc_converter import cli


class ProtocolTests(unittest.TestCase):
    def run_cli(self, args):
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            code = cli.main(args)
        return code, output.getvalue()

    def test_content_modes_keep_output_complete(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / 'large.docx'
            doc = Document()
            doc.add_paragraph('正文😀' * 10000)
            doc.save(source)
            for mode in ('none', 'preview', 'full'):
                code, stdout = self.run_cli(['convert', str(source), '--content', mode, '--preview-chars', '8'])
                result = json.loads(stdout)
                self.assertEqual(0, code, result)
                complete = Path(result['output_path']).read_text(encoding='utf-8')
                self.assertEqual(len(complete), result['markdown_chars'])
                if mode == 'none':
                    self.assertNotIn('markdown_content', result)
                    self.assertLess(len(stdout), 1500)
                else:
                    self.assertEqual(complete[:8] if mode == 'preview' else complete, result['markdown_content'])

    def test_batch_cap_manifest_and_partial_failure(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / 'a.docx'
            doc = Document(); doc.add_paragraph('body'); doc.save(source)
            (Path(td) / 'b.docx').write_bytes(b'broken')
            code, stdout = self.run_cli(['batch', td, '--content', 'none', '--manifest', 'auto', '--max-results', '1'])
            result = json.loads(stdout)
            self.assertEqual(1, code)
            self.assertEqual(2, result['total'])
            self.assertEqual(1, result['succeeded'])
            self.assertEqual(1, result['failed'])
            self.assertEqual(1, result['omitted'])
            self.assertEqual(1, len(result['results']))
            records = [json.loads(line) for line in Path(result['manifest_path']).read_text(encoding='utf-8').splitlines()]
            self.assertEqual(['result', 'result', 'summary'], [r['type'] for r in records])
            self.assertNotIn('markdown_content', json.dumps(records))

    def test_jsonl_flushes_before_requesting_next_file(self):
        with tempfile.TemporaryDirectory() as td:
            manifest = Path(td) / 'manifest.jsonl'
            def conversions(*args, **kwargs):
                yield {'file': 'one.docx', 'result': {'success': True, 'output_path': 'one.md', 'markdown_content': 'first'}}
                persisted = manifest.read_text(encoding='utf-8')
                self.assertEqual('result', json.loads(persisted)['type'])
                yield {'file': 'two.docx', 'result': {'success': False, 'error_code': 'CONVERSION_ERROR', 'error': 'broken'}}
            with patch('bruce_doc_converter.cli.iter_batch_convert', side_effect=conversions):
                code, stdout = self.run_cli(['batch', td, '--jsonl', '--manifest', str(manifest)])
            records = [json.loads(line) for line in stdout.splitlines()]
            self.assertEqual(1, code)
            self.assertEqual(['result', 'result', 'summary'], [r['type'] for r in records])
            self.assertEqual(1, records[-1]['succeeded'])

    def test_manifest_never_overwrites_an_input(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / 'source.docx'
            source.write_bytes(b'original')
            code, stdout = self.run_cli(['batch', td, '--manifest', str(source)])
            self.assertEqual(1, code)
            self.assertEqual('BATCH_IO_ERROR', json.loads(stdout)['error_code'])
            self.assertEqual(b'original', source.read_bytes())

    def test_strict_cli_keeps_failure_diagnostics(self):
        result = {'success': False, 'error_code': 'CONTENT_INCOMPLETE', 'error': 'incomplete',
                  'warnings': ['missing cache'], 'diagnostics': [{'code': 'FORMULA_CACHE_MISSING', 'severity': 'warning', 'message': 'missing cache', 'cell': 'B2'}]}
        with patch('bruce_doc_converter.cli.convert_document', return_value=result) as convert:
            code, stdout = self.run_cli(['convert', 'sheet.xlsx', '--strict'])
        self.assertTrue(convert.call_args.kwargs['strict'])
        self.assertEqual(1, code)
        self.assertEqual(result['diagnostics'], json.loads(stdout)['diagnostics'])

    def test_capping_requires_a_manifest(self):
        code, stdout = self.run_cli(['batch', '.', '--max-results', '1'])
        self.assertEqual(1, code)
        self.assertEqual('USAGE_ERROR', json.loads(stdout)['error_code'])
