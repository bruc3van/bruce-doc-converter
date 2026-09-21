import hashlib
import io
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from openpyxl import Workbook
from PIL import Image

from bruce_doc_converter.converter import convert_document


class IntegrityTests(unittest.TestCase):
    def make_docx(self, path, color):
        image = io.BytesIO()
        Image.new('RGB', (200, 150), color).save(image, format='PNG', compress_level=0)
        image.seek(0)
        doc = Document()
        doc.add_paragraph(color)
        doc.add_picture(image)
        doc.save(path)

    def test_repeated_conversion_preserves_old_images(self):
        with tempfile.TemporaryDirectory() as td:
            source, out = Path(td) / 'report.docx', Path(td) / 'out'
            self.make_docx(source, 'red')
            first = convert_document(source, output_dir=out)
            image = out / first['extracted_images'][0]
            digest = hashlib.sha256(image.read_bytes()).digest()
            self.make_docx(source, 'blue')
            second = convert_document(source, output_dir=out)
            self.assertTrue(second['success'], second)
            self.assertNotEqual(first['extracted_images'], second['extracted_images'])
            self.assertEqual(digest, hashlib.sha256(image.read_bytes()).digest())

    def test_concurrent_outputs_are_independent(self):
        with tempfile.TemporaryDirectory() as td:
            source, out = Path(td) / 'report.docx', Path(td) / 'out'
            self.make_docx(source, 'red')
            with ThreadPoolExecutor(max_workers=4) as pool:
                results = list(pool.map(lambda _: convert_document(source, output_dir=out), range(8)))
            self.assertTrue(all(r['success'] for r in results), results)
            self.assertEqual(8, len({r['output_path'] for r in results}))
            self.assertEqual(8, len({r['extracted_images'][0] for r in results}))
            for r in results:
                self.assertEqual(r['markdown_content'], Path(r['output_path']).read_text(encoding='utf-8'))

    def test_uncached_formula_is_preserved_with_location(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / 'formula.xlsx'
            book = Workbook()
            book.active.append(['item', 'value'])
            book.active.append(['subtotal', '=1+2'])
            book.save(source)
            result = convert_document(source)
            self.assertTrue(result['success'], result)
            self.assertIn('=1+2', result['markdown_content'])
            self.assertEqual('FORMULA_CACHE_MISSING', result['diagnostics'][0]['code'])
            self.assertEqual('B2', result['diagnostics'][0]['cell'])

    def test_failed_write_removes_only_its_reservation_and_images(self):
        with tempfile.TemporaryDirectory() as td:
            source, out = Path(td) / 'report.docx', Path(td) / 'out'
            self.make_docx(source, 'red')
            first = convert_document(source, output_dir=out)
            original = Path(first['output_path']).read_bytes()
            before = {p.relative_to(out) for p in out.rglob('*')}
            with patch('bruce_doc_converter.converter.atomic_write', side_effect=OSError('disk full')):
                second = convert_document(source, output_dir=out)
            self.assertFalse(second['success'])
            self.assertEqual(original, Path(first['output_path']).read_bytes())
            self.assertEqual(before, {p.relative_to(out) for p in out.rglob('*')})

    def test_pdf_partial_failure_is_visible_and_strict_rejects_it(self):
        for strict in (False, True):
            with self.subTest(strict=strict), tempfile.TemporaryDirectory() as td:
                source = Path(td) / 'input.pdf'
                source.write_bytes(b'placeholder')
                pdf, good, bad = MagicMock(), MagicMock(), MagicMock()
                pdf.pages = [good, bad]
                bad.find_tables.side_effect = ValueError('broken page')
                bad.extract_text.side_effect = ValueError('no fallback')
                with patch('pdfplumber.open') as op, patch(
                    'bruce_doc_converter.formats.pdf._extract_pdf_page_blocks', return_value=[(0, 'good page')]
                ):
                    op.return_value.__enter__.return_value = pdf
                    result = convert_document(source, strict=strict)
                self.assertEqual(not strict, result['success'], result)
                self.assertEqual('failed', result['diagnostics'][1]['status'])
                self.assertEqual(2, result['diagnostics'][1]['page'])
                self.assertTrue(result['warnings'])
                if strict:
                    self.assertEqual('CONTENT_INCOMPLETE', result['error_code'])
                    self.assertFalse(list((Path(td) / 'Markdown').glob('*.md')))


if __name__ == '__main__':
    unittest.main()
