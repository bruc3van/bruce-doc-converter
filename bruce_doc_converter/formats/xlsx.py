"""Xlsx conversion helpers."""
from .common import (
    _normalize_table_cell,
    _normalize_text,
    _table_position_has_content,
)
from .images import (
    _image_warning,
    _is_decorative_image,
    _make_image_markdown,
    _save_extracted_image,
)
import os
import re
import logging

logger = logging.getLogger(__name__)


def convert_xlsx(file_path, image_save_dir=None, image_rel_dir=None, diagnostics=None):
    """转换 Excel 文件，支持多表头、空白分隔区、冻结窗格、常见格式保留和图片提取"""
    import openpyxl
    from datetime import date, datetime, time

    workbook = openpyxl.load_workbook(file_path, data_only=True)
    diagnostics = diagnostics if diagnostics is not None else []
    # openpyxl does not calculate formulas. Restore only missing cached results,
    # keeping cached values and their original cell formatting intact.
    try:
        formulas = openpyxl.load_workbook(file_path, data_only=False, read_only=True)
        try:
            for sheet in formulas:
                cached_sheet = workbook[sheet.title]
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type == 'f' and cached_sheet[cell.coordinate].value is None:
                            formula = cell.value
                            if not isinstance(formula, str):
                                formula = getattr(formula, 'text', None) or '[unsupported formula]'
                            cached_sheet[cell.coordinate] = formula
                            diagnostics.append({
                                'code': 'FORMULA_CACHE_MISSING', 'severity': 'warning',
                                'sheet': sheet.title, 'cell': cell.coordinate,
                                'message': f'{sheet.title}!{cell.coordinate}: 公式无缓存结果，已保留公式，未计算。',
                            })
        finally:
            formulas.close()
    except Exception:
        workbook.close()
        raise
    content = ""
    image_counter = 0
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    extracted_images = []

    def _build_merge_map(worksheet):
        """构建合并单元格续格坐标集合，用于保留占位但不重复填充值"""
        merge_map = set()
        for merge_range in worksheet.merged_cells.ranges:
            for row in range(merge_range.min_row, merge_range.max_row + 1):
                for col in range(merge_range.min_col, merge_range.max_col + 1):
                    if row == merge_range.min_row and col == merge_range.min_col:
                        continue
                    merge_map.add((row, col))
        return merge_map

    def _count_number_format_decimals(number_format):
        fmt = (number_format or "").split(";")[0]
        fmt = re.sub(r'"[^"]*"', "", fmt)
        fmt = re.sub(r"\[[^\]]*\]", "", fmt)
        match = re.search(r"\.([0#]+)", fmt)
        return len(match.group(1)) if match else 0

    def _format_excel_datetime(value):
        if isinstance(value, datetime):
            if value.time() == datetime.min.time():
                return value.strftime("%Y-%m-%d")
            return value.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
        return _normalize_text(value)

    def _format_excel_number(value, number_format):
        decimals = _count_number_format_decimals(number_format)
        use_grouping = "," in (number_format or "")

        if "%" in (number_format or ""):
            scaled = value * 100
            formatted = f"{scaled:,.{decimals}f}" if use_grouping else f"{scaled:.{decimals}f}"
            return f"{formatted}%"

        if decimals > 0:
            return f"{value:,.{decimals}f}" if use_grouping else f"{value:.{decimals}f}"

        if isinstance(value, float) and not value.is_integer():
            text = f"{value:,}" if use_grouping else f"{value}"
            return text.rstrip("0").rstrip(".")

        integer_value = int(round(value))
        return f"{integer_value:,}" if use_grouping else str(integer_value)

    def _format_excel_cell(cell, is_merged_placeholder=False):
        value = cell.value
        if value is None:
            return "" if is_merged_placeholder else None

        if cell.is_date or isinstance(value, (datetime, date, time)):
            return _format_excel_datetime(value)

        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"

        if isinstance(value, (int, float)):
            return _format_excel_number(value, cell.number_format or "")

        return _normalize_text(value)

    def _classify_excel_cell(cell, is_merged_placeholder=False):
        if is_merged_placeholder:
            return "placeholder"
        value = cell.value
        if value is None:
            return "blank"
        if cell.is_date or isinstance(value, (datetime, date, time)):
            return "date"
        if isinstance(value, bool):
            return "text"
        if isinstance(value, (int, float)):
            return "number"
        return "text"

    def _iter_table_row_groups(worksheet, merge_map):
        current_group = []
        for row in worksheet.iter_rows(values_only=False):
            cells = []
            occupied_positions = []
            display_values = []
            kinds = []

            for cell in row:
                coord = (cell.row, cell.column)
                is_merged_placeholder = coord in merge_map
                display_value = _format_excel_cell(cell, is_merged_placeholder=is_merged_placeholder)
                occupied = _table_position_has_content(display_value, occupied=is_merged_placeholder)

                cells.append(cell)
                display_values.append(display_value)
                occupied_positions.append(occupied)
                kinds.append(_classify_excel_cell(cell, is_merged_placeholder=is_merged_placeholder))

            row_record = {
                "row_index": row[0].row if row else 0,
                "cells": cells,
                "values": display_values,
                "occupied": occupied_positions,
                "kinds": kinds,
            }

            if any(occupied_positions):
                current_group.append(row_record)
            elif current_group:
                yield current_group
                current_group = []

        if current_group:
            yield current_group

    def _split_column_segments(row_group):
        width = max((len(row["occupied"]) for row in row_group), default=0)
        active_columns = [any(idx < len(row["occupied"]) and row["occupied"][idx] for row in row_group) for idx in range(width)]
        segments = []
        start = None

        for idx, is_active in enumerate(active_columns):
            if is_active and start is None:
                start = idx
            elif not is_active and start is not None:
                segments.append((start, idx))
                start = None

        if start is not None:
            segments.append((start, width))
        return segments

    def _slice_table_rows(row_group, col_start, col_end):
        sliced_rows = []
        for row in row_group:
            values = row["values"][col_start:col_end]
            occupied = row["occupied"][col_start:col_end]
            kinds = row["kinds"][col_start:col_end]
            if not any(occupied):
                continue
            sliced_rows.append({
                "row_index": row["row_index"],
                "values": values,
                "occupied": occupied,
                "kinds": kinds,
            })
        return sliced_rows

    def _profile_table_row(row_data):
        text_count = 0
        number_count = 0
        date_count = 0
        placeholder_count = 0
        non_empty_count = 0

        for value, occupied, kind in zip(row_data["values"], row_data["occupied"], row_data["kinds"]):
            if not occupied:
                continue
            if kind == "placeholder":
                placeholder_count += 1
                continue
            if value not in (None, ""):
                non_empty_count += 1
            if kind == "text":
                text_count += 1
            elif kind == "number":
                number_count += 1
            elif kind == "date":
                date_count += 1

        return {
            "text_count": text_count,
            "number_count": number_count,
            "date_count": date_count,
            "placeholder_count": placeholder_count,
            "non_empty_count": non_empty_count,
            "looks_header": non_empty_count > 0 and text_count >= (number_count + date_count),
            "looks_data": (number_count + date_count) > text_count,
        }

    def _get_freeze_header_rows(worksheet):
        freeze_panes = worksheet.freeze_panes
        if not freeze_panes:
            return 0
        if hasattr(freeze_panes, "row"):
            return max(int(freeze_panes.row) - 1, 0)
        match = re.match(r"[A-Za-z]+(\d+)", str(freeze_panes))
        if not match:
            return 0
        return max(int(match.group(1)) - 1, 0)

    def _determine_header_row_count(table_rows, freeze_header_rows):
        if not table_rows:
            return 0

        frozen_header_count = 0
        if freeze_header_rows > 0:
            for row in table_rows:
                if row["row_index"] <= freeze_header_rows:
                    frozen_header_count += 1
                else:
                    break
            if frozen_header_count > 0:
                return min(frozen_header_count, len(table_rows))

        first_profile = _profile_table_row(table_rows[0])
        if not first_profile["looks_header"]:
            return 0

        header_count = 1
        if len(table_rows) >= 3:
            second_profile = _profile_table_row(table_rows[1])
            third_profile = _profile_table_row(table_rows[2])
            if first_profile["placeholder_count"] > 0 and second_profile["looks_header"] and third_profile["looks_data"]:
                header_count = 2

        return min(header_count, len(table_rows))

    def _build_header_labels(header_rows, col_count):
        if not header_rows:
            return [f"Column {idx + 1}" for idx in range(col_count)]

        expanded_rows = []
        for row in header_rows:
            carry_text = ""
            expanded = []
            for value, occupied in zip(row["values"], row["occupied"]):
                text = _normalize_text(value)
                if text:
                    carry_text = text
                    expanded.append(text)
                elif len(header_rows) > 1 and occupied and carry_text:
                    expanded.append(carry_text)
                else:
                    expanded.append("")
            expanded_rows.append(expanded)

        headers = []
        for col_idx in range(col_count):
            parts = []
            for row in expanded_rows:
                if col_idx >= len(row):
                    continue
                part = row[col_idx].strip()
                if part and (not parts or parts[-1] != part):
                    parts.append(part)
            headers.append(" / ".join(parts) if parts else "")
        return headers

    def _render_table_block(table_rows, freeze_header_rows):
        col_count = max((len(row["values"]) for row in table_rows), default=0)
        if col_count == 0:
            return ""

        header_count = _determine_header_row_count(table_rows, freeze_header_rows)
        header_rows = table_rows[:header_count]
        headers = _build_header_labels(header_rows, col_count)
        data_rows = table_rows[header_count:] if header_count > 0 else table_rows

        lines = [
            "| " + " | ".join(_normalize_table_cell(header) for header in headers) + " |",
            "| " + " | ".join(["---"] * col_count) + " |",
        ]

        for row in data_rows:
            row_values = []
            for idx in range(col_count):
                value = row["values"][idx] if idx < len(row["values"]) else ""
                occupied = row["occupied"][idx] if idx < len(row["occupied"]) else False
                if _table_position_has_content(value, occupied):
                    row_values.append(_normalize_table_cell(value))
                else:
                    row_values.append("")
            lines.append("| " + " | ".join(row_values) + " |")

        return "\n".join(lines)

    try:
        for sheet_name in workbook.sheetnames:
            if len(workbook.sheetnames) > 1:
                content += f"## {_normalize_text(sheet_name)}\n\n"

            worksheet = workbook[sheet_name]
            merge_map = _build_merge_map(worksheet)
            freeze_header_rows = _get_freeze_header_rows(worksheet)
            table_blocks = []

            for row_group in _iter_table_row_groups(worksheet, merge_map):
                for col_start, col_end in _split_column_segments(row_group):
                    table_rows = _slice_table_rows(row_group, col_start, col_end)
                    if not table_rows:
                        continue
                    table_markdown = _render_table_block(table_rows, freeze_header_rows)
                    if table_markdown:
                        table_blocks.append(table_markdown)

            if len(table_blocks) == 1:
                content += table_blocks[0] + "\n\n"
            elif len(table_blocks) > 1:
                for idx, table_markdown in enumerate(table_blocks, 1):
                    content += f"### Table {idx}\n\n{table_markdown}\n\n"

            # 提取 worksheet 中的嵌入图片
            if image_save_dir is not None:
                try:
                    ws_images = getattr(worksheet, '_images', []) or []
                    # 按锚定行号排序
                    sorted_images = []
                    for img_obj in ws_images:
                        try:
                            anchor = getattr(img_obj, 'anchor', None)
                            anchor_from = getattr(anchor, '_from', None) if anchor else None
                            row = getattr(anchor_from, 'row', 0) if anchor_from else 0
                            col = getattr(anchor_from, 'col', 0) if anchor_from else 0
                            sorted_images.append((row, col, img_obj))
                        except Exception:
                            logger.debug("Failed to read XLSX image anchor; falling back to default order", exc_info=True)
                            sorted_images.append((0, 0, img_obj))
                    sorted_images.sort(key=lambda x: (x[0], x[1]))

                    for _row, _col, img_obj in sorted_images:
                        try:
                            # 获取图片数据
                            img_ref = getattr(img_obj, 'ref', None) or getattr(img_obj, '_data', None)
                            image_data = None
                            if img_ref is not None:
                                # openpyxl Image 对象的图片数据
                                if hasattr(img_ref, 'read'):
                                    img_ref.seek(0)
                                    image_data = img_ref.read()
                                elif isinstance(img_ref, bytes):
                                    image_data = img_ref
                            # 回退：尝试从 _data 属性读取
                            if image_data is None and hasattr(img_obj, '_data'):
                                raw = img_obj._data
                                if callable(raw):
                                    raw = raw()
                                if isinstance(raw, bytes):
                                    image_data = raw
                                elif hasattr(raw, 'read'):
                                    raw.seek(0)
                                    image_data = raw.read()

                            if not image_data:
                                _image_warning(diagnostics, sheet_name, '图片数据为空。', sheet=sheet_name,
                                               cell=f'{openpyxl.utils.get_column_letter(_col + 1)}{_row + 1}')
                                continue

                            # 装饰性过滤
                            if _is_decorative_image(image_data):
                                continue

                            image_counter += 1
                            rel_path = _save_extracted_image(
                                image_data, image_save_dir, image_rel_dir,
                                base_name, image_counter
                            )
                            if rel_path:
                                extracted_images.append(rel_path)
                                content += f"{_make_image_markdown(rel_path)}\n\n"
                            else:
                                _image_warning(diagnostics, sheet_name, '无法保存提取的图片。', sheet=sheet_name,
                                               cell=f'{openpyxl.utils.get_column_letter(_col + 1)}{_row + 1}')
                        except Exception as exc:
                            logger.debug("Failed to extract an XLSX embedded image; skipping it", exc_info=True)
                            _image_warning(diagnostics, sheet_name, str(exc), sheet=sheet_name)
                            continue
                except Exception as exc:
                    logger.debug("Failed to inspect XLSX worksheet images; continuing without them", exc_info=True)
                    _image_warning(diagnostics, sheet_name, str(exc), sheet=sheet_name)
    finally:
        workbook.close()

    return content.strip(), extracted_images
