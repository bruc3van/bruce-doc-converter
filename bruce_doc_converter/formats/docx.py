"""Docx conversion helpers."""
from .common import (
    DOCX_W_NS,
    DOCX_XML_NAMESPACES,
    OOXML_IMAGE_NAMESPACES,
    _compose_inline_markdown,
    _normalize_table_cell,
    _normalize_text,
)
from .images import (
    _image_warning,
    _check_ooxml_decorative_flag,
    _is_decorative_image,
    _make_image_markdown,
    _save_extracted_image,
)
import os
import re
import xml.etree.ElementTree as ET
import logging

logger = logging.getLogger(__name__)


def _resolve_docx_style_font_flag(style, attr_name):
    """沿样式继承链解析 tri-state 字体属性，返回 True/False/None"""
    visited = set()
    current = style
    while current is not None:
        style_id = id(current)
        if style_id in visited:
            break
        visited.add(style_id)

        font = getattr(current, "font", None)
        value = getattr(font, attr_name, None) if font is not None else None
        if value is not None:
            return bool(value)

        current = getattr(current, "base_style", None)

    return None


def _get_docx_heading_level(style):
    """根据 Word 样式解析标题层级，无法识别时返回 None"""
    if style is None:
        return None

    style_id = getattr(style, "style_id", "")
    if isinstance(style_id, str) and style_id:
        match = re.match(r"(?i)^heading(\d+)$", style_id.strip())
        if match:
            return int(match.group(1))

    style_name = getattr(style, "name", "")
    if isinstance(style_name, str) and style_name:
        match = re.match(r"^(?:Heading|标题)\s*(\d+)$", style_name.strip())
        if match:
            return int(match.group(1))

    return None


def _resolve_docx_run_font_flag(run, paragraph, attr_name, *, allow_paragraph_style=False):
    """解析 run 的实际粗体/斜体状态，支持字符样式和可选的段落样式继承"""
    direct_value = getattr(run.font, attr_name, None)
    if direct_value is not None:
        return bool(direct_value)

    char_style_value = _resolve_docx_style_font_flag(getattr(run, "style", None), attr_name)
    if char_style_value is not None:
        return char_style_value

    if allow_paragraph_style:
        paragraph_style_value = _resolve_docx_style_font_flag(getattr(paragraph, "style", None), attr_name)
        if paragraph_style_value is not None:
            return paragraph_style_value

    return False


def _get_docx_grid_span(tc):
    """读取 Word 表格单元格的横向跨列数"""
    tc_pr = getattr(tc, "tcPr", None)
    grid_span = getattr(tc_pr, "gridSpan", None) if tc_pr is not None else None
    raw_value = getattr(grid_span, "val", None) if grid_span is not None else None
    try:
        return max(int(raw_value), 1)
    except (TypeError, ValueError):
        return 1


def _is_docx_vertical_merge_continuation(tc):
    """判断 Word 表格单元格是否为纵向合并的续格"""
    tc_pr = getattr(tc, "tcPr", None)
    if tc_pr is None:
        return False

    v_merge = getattr(tc_pr, "vMerge", None)
    if v_merge is None:
        return False

    return getattr(v_merge, "val", None) != "restart"


def _extract_docx_table_cell_text(tc):
    """从 Word 表格 XML 单元格中提取文本，保留段落换行"""
    paragraphs = []
    try:
        root = ET.fromstring(tc.xml)
    except ET.ParseError:
        return ""

    for paragraph in root.findall('./w:p', DOCX_XML_NAMESPACES):
        text_nodes = paragraph.findall('.//w:t', DOCX_XML_NAMESPACES)
        paragraph_text = ''.join(node.text for node in text_nodes if node.text)
        if paragraph_text:
            paragraphs.append(paragraph_text)

    if paragraphs:
        return _normalize_table_cell("\n".join(paragraphs))

    text_nodes = root.findall('.//w:t', DOCX_XML_NAMESPACES)
    return _normalize_table_cell(''.join(node.text for node in text_nodes if node.text))


def _docx_attr(node, attr_name):
    """读取 WordprocessingML 命名空间属性值"""
    if node is None:
        return None
    return node.attrib.get(f'{{{DOCX_W_NS}}}{attr_name}')


def _build_docx_numbering_index(doc):
    """构建 numId -> 抽象编号定义 的索引，支持多级编号渲染"""
    try:
        root = ET.fromstring(doc.part.numbering_part.element.xml)
    except (AttributeError, ET.ParseError, TypeError):
        return {}, {}

    num_to_abstract = {}
    abstract_levels = {}

    for num_node in root.findall('.//w:num', DOCX_XML_NAMESPACES):
        num_id = _docx_attr(num_node, 'numId')
        abstract_node = num_node.find('./w:abstractNumId', DOCX_XML_NAMESPACES)
        abstract_id = _docx_attr(abstract_node, 'val')
        if num_id and abstract_id:
            num_to_abstract[str(num_id)] = str(abstract_id)

    for abstract_node in root.findall('.//w:abstractNum', DOCX_XML_NAMESPACES):
        abstract_id = _docx_attr(abstract_node, 'abstractNumId')
        if not abstract_id:
            continue

        levels = {}
        for level_node in abstract_node.findall('./w:lvl', DOCX_XML_NAMESPACES):
            ilvl_raw = _docx_attr(level_node, 'ilvl')
            try:
                ilvl = int(ilvl_raw)
            except (TypeError, ValueError):
                continue

            start_node = level_node.find('./w:start', DOCX_XML_NAMESPACES)
            num_fmt_node = level_node.find('./w:numFmt', DOCX_XML_NAMESPACES)
            lvl_text_node = level_node.find('./w:lvlText', DOCX_XML_NAMESPACES)

            try:
                start = int(_docx_attr(start_node, 'val') or 1)
            except (TypeError, ValueError):
                start = 1

            levels[ilvl] = {
                'start': start,
                'num_fmt': _docx_attr(num_fmt_node, 'val') or '',
                'lvl_text': _docx_attr(lvl_text_node, 'val') or f'%{ilvl + 1}.',
            }

        if levels:
            abstract_levels[str(abstract_id)] = levels

    return num_to_abstract, abstract_levels


def _get_docx_style_numpr(style):
    """从段落样式中读取继承的 numPr，兼容内建 List Number / List Bullet 样式"""
    if style is None or not hasattr(style, 'element'):
        return None, None

    try:
        num_id_nodes = style.element.xpath('./w:pPr/w:numPr/w:numId')
        ilvl_nodes = style.element.xpath('./w:pPr/w:numPr/w:ilvl')
    except Exception:
        return None, None

    num_id = _docx_attr(num_id_nodes[0], 'val') if num_id_nodes else None
    ilvl_raw = _docx_attr(ilvl_nodes[0], 'val') if ilvl_nodes else None

    try:
        ilvl = int(ilvl_raw) if ilvl_raw is not None else 0
    except (TypeError, ValueError):
        ilvl = 0

    return (str(num_id), ilvl) if num_id is not None else (None, None)


def _get_docx_paragraph_numpr(para):
    """获取段落实际使用的 numId / ilvl，优先段落自身，再回退样式"""
    p = getattr(para, "_p", None)
    ppr = getattr(p, "pPr", None) if p is not None else None
    num_pr = getattr(ppr, "numPr", None) if ppr is not None else None

    if num_pr is not None:
        num_id_el = getattr(num_pr, "numId", None)
        ilvl_el = getattr(num_pr, "ilvl", None)
        num_id = getattr(num_id_el, "val", None)
        ilvl_raw = getattr(ilvl_el, "val", None)
        try:
            ilvl = int(ilvl_raw) if ilvl_raw is not None else 0
        except (TypeError, ValueError):
            ilvl = 0
        if num_id is not None:
            return str(num_id), ilvl

    return _get_docx_style_numpr(getattr(para, "style", None))


def _to_roman(value):
    if value <= 0:
        return str(value)
    pairs = [
        (1000, "M"), (900, "CM"), (500, "D"), (400, "CD"),
        (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
        (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"),
    ]
    result = []
    remaining = value
    for number, numeral in pairs:
        while remaining >= number:
            result.append(numeral)
            remaining -= number
    return ''.join(result)


def _to_alpha(value, uppercase=False):
    if value <= 0:
        return str(value)
    letters = []
    current = value
    while current > 0:
        current -= 1
        letters.append(chr((current % 26) + (65 if uppercase else 97)))
        current //= 26
    return ''.join(reversed(letters))


def _to_chinese_counting(value):
    if value <= 0:
        return str(value)

    digits = "零一二三四五六七八九"
    units = ["", "十", "百", "千"]
    if value < 10:
        return digits[value]
    if value < 10000:
        parts = []
        zero_pending = False
        chars = list(str(value))
        length = len(chars)
        for idx, char in enumerate(chars):
            digit = int(char)
            unit = units[length - idx - 1]
            if digit == 0:
                zero_pending = bool(parts)
                continue
            if zero_pending:
                parts.append("零")
                zero_pending = False
            if not (digit == 1 and unit == "十" and not parts):
                parts.append(digits[digit])
            parts.append(unit)
        return ''.join(parts) or digits[0]
    return str(value)


def _to_circled_number(value):
    if 1 <= value <= 20:
        return chr(9311 + value)
    return str(value)


def _format_docx_number_value(value, num_fmt):
    """按 Word numFmt 渲染编号文本"""
    fmt = (num_fmt or '').lower()
    if fmt in {'decimal', 'decimalfullwidth'}:
        return str(value)
    if fmt == 'decimalzero':
        return f"{value:02d}"
    if fmt == 'lowerletter':
        return _to_alpha(value, uppercase=False)
    if fmt == 'upperletter':
        return _to_alpha(value, uppercase=True)
    if fmt == 'lowerroman':
        return _to_roman(value).lower()
    if fmt == 'upperroman':
        return _to_roman(value)
    if fmt in {'chinesecounting', 'chineselegalsimplified', 'ideographtraditional', 'taiwanesecounting'}:
        return _to_chinese_counting(value)
    if fmt in {'decimalenclosedcircle', 'circleNumDbPlain'.lower(), 'decimalenclosedcirclechinese'}:
        return _to_circled_number(value)
    return str(value)


def _render_docx_list_marker(numbering_info, numbering_state):
    """渲染 Word 多级编号列表的 Markdown 前缀"""
    if not numbering_info:
        return None
    if not numbering_info["ordered"]:
        return "-"

    num_id = numbering_info["num_id"]
    level = numbering_info["level"]
    levels = numbering_info["levels"]
    level_def = levels.get(level, {})
    state = numbering_state.setdefault(num_id, {})

    for existing_level in list(state.keys()):
        if existing_level > level:
            del state[existing_level]

    if level not in state:
        state[level] = max(level_def.get("start", 1) - 1, 0)
    state[level] += 1

    for ancestor_level in range(level):
        if ancestor_level not in state:
            ancestor_def = levels.get(ancestor_level, {})
            state[ancestor_level] = max(ancestor_def.get("start", 1), 1)

    template = level_def.get("lvl_text") or f"%{level + 1}."

    def _replace(match):
        ref_level = int(match.group(1)) - 1
        ref_value = state.get(ref_level, 1)
        ref_def = levels.get(ref_level, level_def)
        return _format_docx_number_value(ref_value, ref_def.get("num_fmt"))

    rendered = re.sub(r"%(\d+)", _replace, template).strip()
    return rendered or "1."


def _is_docx_toc_paragraph(para):
    """识别 Word 自动目录段落，避免被误当正文导出"""
    style = getattr(para, "style", None)
    style_name = getattr(style, "name", "") if style is not None else ""
    style_id = getattr(style, "style_id", "") if style is not None else ""

    if re.match(r"(?i)^toc(?:\s+heading|\s+\d+)?$", style_name.strip()):
        return True
    if re.match(r"(?i)^toc(?:heading|\d+)?$", style_id.strip()):
        return True

    try:
        root = ET.fromstring(para._p.xml)
    except Exception:
        return False

    for instr in root.findall('.//w:instrText', DOCX_XML_NAMESPACES):
        if 'TOC' in (instr.text or '').upper():
            return True
    return False


def convert_docx(file_path, image_save_dir=None, image_rel_dir=None, diagnostics=None):
    """转换 Word 文档，支持标题、格式、列表（含编号/层级）和图片提取"""
    import docx

    doc = docx.Document(file_path)
    content = ""
    num_to_abstract, abstract_levels = _build_docx_numbering_index(doc)
    numbering_state = {}
    image_counter = 0
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    extracted_images = []

    diagnostics = diagnostics if diagnostics is not None else []

    def _extract_drawing_images(paragraph_element):
        """
        从段落 XML 中提取图片。
        支持 w:drawing（内联和浮动）以及 mc:AlternateContent 包裹的图片。

        Returns:
            图片 Markdown 字符串列表
        """
        nonlocal image_counter

        if image_save_dir is None:
            return []

        image_markdowns = []
        ns = OOXML_IMAGE_NAMESPACES

        try:
            p_xml = ET.fromstring(paragraph_element.xml)
        except (ET.ParseError, AttributeError):
            return []

        # 查找所有 drawing 元素（直接和通过 mc:AlternateContent 包裹的）
        drawings = []
        # 直接 w:drawing
        for drawing in p_xml.findall('.//w:drawing', ns):
            drawings.append(drawing)
        # mc:AlternateContent -> mc:Choice -> w:drawing
        for alt_content in p_xml.findall('.//mc:AlternateContent', ns):
            for choice in alt_content.findall('.//mc:Choice', ns):
                for drawing in choice.findall('.//w:drawing', ns):
                    if drawing not in drawings:
                        drawings.append(drawing)
            # mc:Fallback 中也可能有图片
            for fallback in alt_content.findall('.//mc:Fallback', ns):
                for drawing in fallback.findall('.//w:drawing', ns):
                    if drawing not in drawings:
                        drawings.append(drawing)

        for drawing in drawings:
            # 获取 docPr 以检查装饰性标记和 alt text
            doc_pr = drawing.find('.//wp:docPr', ns)
            if doc_pr is None:
                doc_pr = drawing.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')

            is_decorative = False
            alt_text = ""
            if doc_pr is not None:
                is_decorative, alt_text = _check_ooxml_decorative_flag(doc_pr, ns)
            if is_decorative:
                continue

            # 获取图片数据：通过 a:blip 的 r:embed 属性
            blip = drawing.find('.//a:blip', ns)
            if blip is None:
                continue

            embed_id = blip.get(f'{{{ns["r"]}}}embed') or blip.get('embed')
            if not embed_id:
                _image_warning(diagnostics, 'DOCX drawing', '缺少内嵌图片关系，无法提取外部图片。')
                continue

            # 从 document part 的 related_parts 获取图片数据
            try:
                image_part = doc.part.related_parts.get(embed_id)
                if image_part is None:
                    _image_warning(diagnostics, f'DOCX {embed_id}', '图片关系缺失。')
                    continue
                image_data = image_part.blob
            except Exception as exc:
                logger.debug("Failed to read DOCX image part: %s", embed_id, exc_info=True)
                _image_warning(diagnostics, f'DOCX {embed_id}', str(exc))
                continue

            if not image_data:
                _image_warning(diagnostics, f'DOCX {embed_id}', '图片数据为空。')
                continue

            # 装饰性过滤
            if _is_decorative_image(image_data, is_decorative_flag=is_decorative):
                continue

            # 保存图片
            image_counter += 1
            rel_path = _save_extracted_image(
                image_data, image_save_dir, image_rel_dir,
                base_name, image_counter
            )
            if rel_path:
                extracted_images.append(rel_path)
                md = _make_image_markdown(rel_path, alt_text)
                image_markdowns.append(md)
            else:
                _image_warning(diagnostics, f'DOCX {embed_id}', '无法保存提取的图片。')

        return image_markdowns

    def get_numbering_info(para):
        """
        尝试从段落的 numPr / numbering.xml 解析列表信息

        Returns:
            None 或 {'level': int, 'ordered': bool}
        """
        try:
            num_id, level = _get_docx_paragraph_numpr(para)
            if num_id is None:
                return None

            abstract_id = num_to_abstract.get(str(num_id))
            levels = abstract_levels.get(abstract_id, {}) if abstract_id is not None else {}
            level_def = levels.get(level) or levels.get(0) or {}
            num_fmt = level_def.get('num_fmt')

            style = getattr(para, "style", None)
            style_name = getattr(style, "name", "") if style is not None else ""
            style_id = getattr(style, "style_id", "") if style is not None else ""
            style_hint = f"{style_name} {style_id}".lower()

            if (num_fmt or "").lower() == "bullet":
                ordered = False
            elif (num_fmt or ""):
                ordered = True
            elif "bullet" in style_hint or "项目符号" in style_hint or "符号" in style_hint:
                ordered = False
            elif "number" in style_hint or "编号" in style_hint:
                ordered = True
            else:
                ordered = True

            return {
                'level': max(level, 0),
                'ordered': ordered,
                'num_id': str(num_id),
                'levels': levels,
            }
        except Exception:
            return None

    def process_paragraph(para):
        """处理单个段落，识别标题、列表和格式"""
        if _is_docx_toc_paragraph(para):
            return ""
        if not para.text.strip():
            return ""

        style = para.style if hasattr(para, "style") else None
        style_name = style.name if style else ""
        style_id = getattr(style, "style_id", "") if style else ""
        heading_level = _get_docx_heading_level(style)
        allow_paragraph_style = heading_level is None

        # 先拼接富文本（列表项也需要保留粗体/斜体）
        # 将相邻同格式的 run 合并后再添加 Markdown 标记，避免 **text1****text2** 碎片
        groups = []
        for run in para.runs:
            text = run.text
            if not text:
                continue
            fmt = (
                _resolve_docx_run_font_flag(run, para, "bold", allow_paragraph_style=allow_paragraph_style),
                _resolve_docx_run_font_flag(run, para, "italic", allow_paragraph_style=allow_paragraph_style),
            )
            if groups and groups[-1][0] == fmt:
                groups[-1] = (fmt, groups[-1][1] + text)
            else:
                groups.append((fmt, text))
        formatted_text = _compose_inline_markdown(groups)
        text_value = _normalize_text(formatted_text.strip() or para.text.strip())
        if not text_value:
            return ""

        # 识别标题层级
        if heading_level is not None:
            heading_prefix = "#" * min(heading_level, 6)  # Markdown最多支持6级标题
            return f"{heading_prefix} {text_value}\n\n"

        # 检查是否是列表项
        # 优先使用 numPr + numbering.xml 解析列表编号格式与层级
        numbering_info = get_numbering_info(para)
        if numbering_info:
            indent = "    " * numbering_info["level"]
            marker = _render_docx_list_marker(numbering_info, numbering_state)
            return f"{indent}{marker} {text_value}\n"

        # 注意：python-docx对列表的支持有限，这里做基本处理（按样式兜底）
        style_name_str = style_name or ""
        style_id_str = style_id or ""
        if (isinstance(style_id_str, str) and style_id_str.startswith("List")) or (
            isinstance(style_name_str, str) and style_name_str.startswith("List")
        ):
            level = 0
            m = re.search(r"(\d+)$", style_name_str.strip())
            if m:
                level = max(int(m.group(1)) - 1, 0)
            indent = "    " * level
            if "Bullet" in style_id_str or "Bullet" in style_name_str or style_name_str == "List Bullet":
                return f"{indent}- {text_value}\n"
            if "Number" in style_id_str or "Number" in style_name_str or style_name_str == "List Number":
                return f"{indent}1. {text_value}\n"

        return text_value + "\n\n"

    # 处理文档中的所有元素（段落和表格）
    # 需要按照它们在文档中的顺序处理
    paragraphs_iter = iter(doc.paragraphs)
    tables_iter = iter(doc.tables)
    for element in doc.element.body:
        # 处理段落
        if element.tag.endswith('p'):
            para = next(paragraphs_iter, None)
            if para is not None:
                content += process_paragraph(para)
                # 提取段落中的图片
                for img_md in _extract_drawing_images(para._p):
                    content += f"\n{img_md}\n\n"

        # 处理表格
        elif element.tag.endswith('tbl'):
            table = next(tables_iter, None)
            if table is not None:
                # 使用底层 XML 读取真实网格，避免 python-docx 将合并单元格重复展开
                all_rows_data = []
                table_grid = getattr(getattr(table._tbl, "tblGrid", None), "gridCol_lst", None)
                max_cols = len(table_grid) if table_grid is not None else 0

                for tr in table._tbl.tr_lst:
                    row_data = []
                    for tc in tr.tc_lst:
                        span = _get_docx_grid_span(tc)
                        cell_text = "" if _is_docx_vertical_merge_continuation(tc) else _extract_docx_table_cell_text(tc)
                        row_data.append(cell_text)
                        if span > 1:
                            row_data.extend([""] * (span - 1))
                    all_rows_data.append(row_data)

                if not max_cols:
                    max_cols = max((len(r) for r in all_rows_data), default=0)
                if max_cols == 0:
                    continue
                for i, row_data in enumerate(all_rows_data):
                    padded = row_data + [""] * (max_cols - len(row_data))
                    content += "| " + " | ".join(padded) + " |\n"
                    if i == 0:
                        content += "| " + " | ".join(["---"] * max_cols) + " |\n"
                content += "\n"

    return content.strip(), extracted_images
