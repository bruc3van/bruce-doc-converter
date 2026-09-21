"""Pptx conversion helpers."""
from .common import (
    PPTX_BACKGROUND_COVERAGE_RATIO,
    _compose_inline_markdown,
    _escape_plain_markdown_text,
    _normalize_table_cell,
    _normalize_text,
)
from .images import (
    _image_warning,
    _check_ooxml_decorative_flag,
    _is_decorative_image,
    _make_image_markdown,
    _save_extracted_image,
)
import os
import xml.etree.ElementTree as ET
import logging

logger = logging.getLogger(__name__)


def convert_pptx(file_path, image_save_dir=None, image_rel_dir=None, diagnostics=None):
    """转换 PowerPoint 文件，提取标题、正文、表格、图表、图片和备注"""
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    presentation = pptx.Presentation(file_path)
    content = ""
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    image_counter = 0
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    extracted_images = []
    diagnostics = diagnostics if diagnostics is not None else []

    def _resolve_pptx_run_font_flag(run, paragraph, attr_name, *, allow_paragraph_style=False):
        """解析 run 的实际粗体/斜体状态，支持段落默认格式回退"""
        direct_value = getattr(run.font, attr_name, None)
        if direct_value is not None:
            return bool(direct_value)

        if allow_paragraph_style:
            paragraph_value = getattr(paragraph.font, attr_name, None)
            if paragraph_value is not None:
                return bool(paragraph_value)

        return False

    def _process_text_frame(text_frame, role="body"):
        """处理文本框，保留段落层级和格式"""
        result = ""
        for para in text_frame.paragraphs:
            if not para.text.strip():
                continue

            allow_paragraph_style = role != "title"

            # 将相邻同格式的 run 合并后再添加 Markdown 标记，避免 **text1****text2** 碎片
            groups = []
            for run in para.runs:
                text = run.text
                if not text:
                    continue
                fmt = (
                    _resolve_pptx_run_font_flag(run, para, "bold", allow_paragraph_style=allow_paragraph_style),
                    _resolve_pptx_run_font_flag(run, para, "italic", allow_paragraph_style=allow_paragraph_style),
                )
                if groups and groups[-1][0] == fmt:
                    groups[-1] = (fmt, groups[-1][1] + text)
                else:
                    groups.append((fmt, text))
            formatted = _compose_inline_markdown(groups)
            text_value = _normalize_text(formatted.strip() or para.text.strip())
            if not text_value:
                continue

            if role == "title":
                result += f"### {text_value}\n\n"
                continue

            if role == "subtitle":
                result += _escape_plain_markdown_text(text_value) + "\n\n"
                continue

            # 检查列表层级
            level = para.level if para.level else 0
            if level > 0:
                indent = "  " * level
                result += f"{indent}- {text_value}\n"
            elif hasattr(para, '_pPr') and para._pPr is not None and para._pPr.find(
                './/{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None:
                result += f"- {text_value}\n"
            elif hasattr(para, '_pPr') and para._pPr is not None and para._pPr.find(
                './/{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum') is not None:
                result += f"1. {text_value}\n"
            else:
                result += text_value + "\n\n"

        return result

    def _iter_shapes(shapes):
        for shape in shapes:
            yield shape
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                yield from _iter_shapes(shape.shapes)

    def _shape_is_title(shape):
        if not getattr(shape, "is_placeholder", False):
            return False

        try:
            placeholder_type = shape.placeholder_format.type
            return placeholder_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        except Exception:
            return False

    def _get_placeholder_type(shape):
        if not getattr(shape, "is_placeholder", False):
            return None
        try:
            return shape.placeholder_format.type
        except Exception:
            return None

    def _shape_bounds(shape):
        left = getattr(shape, "left", 0)
        top = getattr(shape, "top", 0)
        width = getattr(shape, "width", 0)
        height = getattr(shape, "height", 0)
        return {
            "left": left,
            "top": top,
            "width": width,
            "height": height,
            "right": left + width,
            "bottom": top + height,
            "center_x": left + width / 2,
        }

    def _shape_sort_key(entry):
        return (entry["top"], entry["left"], entry["height"], entry["width"])

    def _render_table_markdown(table):
        all_rows_data = []
        for row in table.rows:
            row_data = []
            seen_cells = set()
            for cell in row.cells:
                cell_id = id(cell)
                if cell_id in seen_cells:
                    continue
                seen_cells.add(cell_id)
                row_data.append(_normalize_table_cell(cell.text))
            all_rows_data.append(row_data)

        max_cols = max((len(r) for r in all_rows_data), default=0)
        if max_cols == 0:
            return ""

        content_part = ""
        for idx, row_data in enumerate(all_rows_data):
            padded = row_data + [""] * (max_cols - len(row_data))
            content_part += "| " + " | ".join(padded) + " |\n"
            if idx == 0:
                content_part += "| " + " | ".join(["---"] * max_cols) + " |\n"
        return content_part.strip()

    def _render_chart_markdown(chart):
        lines = []
        chart_title = ""
        if getattr(chart, "has_title", False):
            try:
                chart_title = _normalize_text(chart.chart_title.text_frame.text)
            except Exception:
                chart_title = ""
        lines.append(f"**Chart:** {chart_title or 'Untitled chart'}")

        series_names = []
        try:
            series_names = [_normalize_text(series.name) for series in chart.series if _normalize_text(series.name)]
        except Exception:
            series_names = []
        if series_names:
            lines.append(f"Series: {', '.join(series_names)}")

        categories = []
        try:
            categories = [_normalize_text(str(category)) for category in chart.plots[0].categories if _normalize_text(str(category))]
        except Exception:
            categories = []
        if categories:
            lines.append(f"Categories: {', '.join(categories)}")

        return "\n".join(lines).strip()

    def _render_picture_markdown(caption_text=None, image_path=None, alt_text=None):
        if image_path:
            alt = alt_text or caption_text or "image"
            md = _make_image_markdown(image_path, alt)
            if caption_text:
                return f"{md}\nCaption: {caption_text}"
            return md
        if caption_text:
            return f"**Image**\nCaption: {caption_text}"
        return "**Image**"

    def _render_diagram_markdown(shape):
        text_value = ""
        if getattr(shape, "has_text_frame", False) and getattr(shape, "text", "").strip():
            text_value = _normalize_text(shape.text)
        shape_name = _normalize_text(getattr(shape, "name", "SmartArt"))
        if text_value:
            return f"**SmartArt:** {shape_name}\n{text_value}"
        return f"**SmartArt:** {shape_name}"

    def _looks_like_title_candidate(entry):
        return (
            entry["kind"] == "text"
            and entry["top"] <= slide_height * 0.22
            and entry["width"] >= slide_width * 0.35
            and len(entry.get("raw_text", "")) <= 120
        )

    def _looks_like_subtitle_candidate(entry, title_entry):
        return (
            entry["kind"] == "text"
            and entry["top"] >= title_entry["bottom"]
            and entry["top"] <= title_entry["bottom"] + slide_height * 0.18
            and entry["width"] >= slide_width * 0.25
            and entry["center_x"] >= slide_width * 0.25
            and entry["center_x"] <= slide_width * 0.75
        )

    def _is_footer_candidate(entry):
        return (
            entry["kind"] == "text"
            and entry["bottom"] >= slide_height * 0.86
            and entry["height"] <= slide_height * 0.12
        )

    def _find_picture_caption(entries, picture_entry):
        best_entry = None
        best_distance = None
        for entry in entries:
            if entry["kind"] != "text" or entry.get("role") != "body" or entry.get("consumed"):
                continue

            horizontal_overlap = min(entry["right"], picture_entry["right"]) - max(entry["left"], picture_entry["left"])
            if horizontal_overlap <= 0:
                continue

            distance = entry["top"] - picture_entry["bottom"]
            if distance < 0 or distance > slide_height * 0.08:
                continue
            if len(entry.get("raw_text", "")) > 160:
                continue

            if best_distance is None or distance < best_distance:
                best_distance = distance
                best_entry = entry
        return best_entry

    def _render_body_entries(entries):
        usable_entries = [entry for entry in entries if entry.get("markdown")]
        if not usable_entries:
            return ""

        left_entries = []
        right_entries = []
        wide_entries = []

        for entry in usable_entries:
            if entry["right"] <= slide_width * 0.48:
                left_entries.append(entry)
            elif entry["left"] >= slide_width * 0.52:
                right_entries.append(entry)
            else:
                wide_entries.append(entry)

        parts = []
        parts.extend(entry["markdown"] for entry in sorted(wide_entries, key=_shape_sort_key))

        if left_entries and right_entries:
            parts.append("#### Left Column")
            parts.extend(entry["markdown"] for entry in sorted(left_entries, key=_shape_sort_key))
            parts.append("#### Right Column")
            parts.extend(entry["markdown"] for entry in sorted(right_entries, key=_shape_sort_key))
            return "\n\n".join(part.strip() for part in parts if part and part.strip()).strip()

        ordered_entries = sorted(usable_entries, key=_shape_sort_key)
        return "\n\n".join(entry["markdown"].strip() for entry in ordered_entries if entry["markdown"].strip()).strip()

    for i, slide in enumerate(presentation.slides, 1):
        slide_parts = []
        if len(presentation.slides) > 1:
            slide_parts.append(f"## Slide {i}")

        entries = []

        for shape in _iter_shapes(slide.shapes):
            bounds = _shape_bounds(shape)
            placeholder_type = _get_placeholder_type(shape)
            entry = {
                "shape": shape,
                "kind": None,
                "role": "body",
                "markdown": "",
                "raw_text": "",
                "placeholder_type": placeholder_type,
                **bounds,
            }

            if getattr(shape, "has_table", False):
                entry["kind"] = "table"
                entry["markdown"] = _render_table_markdown(shape.table)
            elif getattr(shape, "has_chart", False):
                entry["kind"] = "chart"
                entry["markdown"] = _render_chart_markdown(shape.chart)
            elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                entry["kind"] = "picture"
                # 提取图片数据和元数据
                entry["image_path"] = None
                entry["image_alt"] = ""
                if image_save_dir is not None:
                    try:
                        image_data = shape.image.blob
                        if not image_data:
                            raise ValueError('图片数据为空。')
                        # 检查装饰性标记：通过 shape XML 中的 cNvPr
                        is_decorative = False
                        alt_text = ""
                        try:
                            sp_xml = ET.fromstring(shape._element.xml)
                            # PPTX 中 cNvPr 可能在 p:nvPicPr/p:cNvPr 或 nvSpPr/cNvPr
                            cnv_pr = sp_xml.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
                            if cnv_pr is None:
                                cnv_pr = sp_xml.find('.//{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}cNvPr')
                            if cnv_pr is None:
                                # 用更通用的查找
                                for elem in sp_xml.iter():
                                    if elem.tag.endswith('}cNvPr') or elem.tag == 'cNvPr':
                                        cnv_pr = elem
                                        break
                            if cnv_pr is not None:
                                is_decorative, alt_text = _check_ooxml_decorative_flag(cnv_pr)
                        except (AttributeError, ET.ParseError, TypeError):
                            logger.debug("Failed to parse PPTX picture metadata; continuing without decorative flag", exc_info=True)

                        # 检查是否为背景图（覆盖面积 >= 90% 幻灯片）
                        is_background = False
                        if slide_width and slide_height:
                            shape_area = entry["width"] * entry["height"]
                            slide_area = slide_width * slide_height
                            if slide_area > 0 and shape_area / slide_area >= PPTX_BACKGROUND_COVERAGE_RATIO:
                                is_background = True

                        if not _is_decorative_image(
                            image_data,
                            is_decorative_flag=is_decorative,
                            is_pptx_background=is_background
                        ):
                            image_counter += 1
                            rel_path = _save_extracted_image(
                                image_data, image_save_dir, image_rel_dir,
                                base_name, image_counter
                            )
                            if rel_path:
                                extracted_images.append(rel_path)
                                entry["image_path"] = rel_path
                                entry["image_alt"] = alt_text
                            else:
                                _image_warning(diagnostics, f'PPTX slide {i}, shape {shape.shape_id}', '无法保存提取的图片。', page=i)
                    except Exception as exc:
                        logger.debug("Failed to extract a PPTX picture; skipping it", exc_info=True)
                        _image_warning(diagnostics, f'PPTX slide {i}, shape {shape.shape_id}', str(exc), page=i)
            elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.DIAGRAM:
                entry["kind"] = "diagram"
                entry["markdown"] = _render_diagram_markdown(shape)
            elif getattr(shape, "has_text_frame", False) and getattr(shape, "text", "").strip():
                entry["kind"] = "text"
                entry["raw_text"] = _normalize_text(shape.text, preserve_newlines=True)
                if _shape_is_title(shape):
                    entry["role"] = "title"
                elif placeholder_type == PP_PLACEHOLDER.SUBTITLE:
                    entry["role"] = "subtitle"
                elif placeholder_type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER):
                    entry["role"] = "footer"
                entry["markdown"] = _process_text_frame(shape.text_frame, role="title" if entry["role"] == "title" else "body")
            else:
                continue

            if entry["markdown"] or entry["kind"] in {"picture"}:
                entries.append(entry)

        entries.sort(key=_shape_sort_key)

        title_entries = [entry for entry in entries if entry["role"] == "title"]
        subtitle_entries = [entry for entry in entries if entry["role"] == "subtitle"]

        if not title_entries:
            for entry in entries:
                if _looks_like_title_candidate(entry):
                    entry["role"] = "title"
                    entry["markdown"] = _process_text_frame(entry["shape"].text_frame, role="title")
                    title_entries.append(entry)
                    break

        if title_entries and not subtitle_entries:
            title_anchor = sorted(title_entries, key=_shape_sort_key)[0]
            for entry in entries:
                if entry["role"] == "body" and _looks_like_subtitle_candidate(entry, title_anchor):
                    entry["role"] = "subtitle"
                    subtitle_entries.append(entry)
                    break

        for entry in entries:
            if entry["role"] == "body" and _is_footer_candidate(entry):
                entry["role"] = "footer"

        for entry in [entry for entry in entries if entry["kind"] == "picture"]:
            caption_entry = _find_picture_caption(entries, entry)
            caption_text = caption_entry["raw_text"] if caption_entry else None
            if caption_entry is not None:
                caption_entry["consumed"] = True
            entry["markdown"] = _render_picture_markdown(
                caption_text=caption_text,
                image_path=entry.get("image_path"),
                alt_text=entry.get("image_alt")
            )

        title_entries = sorted([entry for entry in entries if entry["role"] == "title"], key=_shape_sort_key)
        subtitle_entries = sorted([entry for entry in entries if entry["role"] == "subtitle"], key=_shape_sort_key)
        footer_entries = sorted([entry for entry in entries if entry["role"] == "footer"], key=_shape_sort_key)
        body_entries = [
            entry for entry in entries
            if entry["role"] == "body" and not entry.get("consumed") and entry["kind"] in {"text", "table"}
        ]
        visual_entries = [
            entry for entry in entries
            if entry["kind"] in {"chart", "picture", "diagram"} and entry.get("markdown")
        ]

        for entry in title_entries:
            if entry["markdown"].strip():
                slide_parts.append(entry["markdown"].strip())

        if subtitle_entries:
            subtitle_body = "\n\n".join(
                _process_text_frame(entry["shape"].text_frame, role="subtitle").strip()
                for entry in subtitle_entries
                if _process_text_frame(entry["shape"].text_frame, role="subtitle").strip()
            ).strip()
            if subtitle_body:
                slide_parts.append("#### Subtitle\n\n" + subtitle_body)

        body_markdown = _render_body_entries(body_entries)
        if body_markdown:
            slide_parts.append(body_markdown)

        if visual_entries:
            visuals_body = "\n\n".join(entry["markdown"].strip() for entry in sorted(visual_entries, key=_shape_sort_key) if entry["markdown"].strip())
            if visuals_body:
                slide_parts.append("#### Visuals\n\n" + visuals_body)

        if footer_entries:
            footer_body = "\n\n".join(
                _process_text_frame(entry["shape"].text_frame, role="subtitle").strip()
                for entry in footer_entries
                if _process_text_frame(entry["shape"].text_frame, role="subtitle").strip()
            ).strip()
            if footer_body:
                slide_parts.append("#### Footer\n\n" + footer_body)

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = _normalize_text(slide.notes_slide.notes_text_frame.text, preserve_newlines=True)
            if notes_text:
                slide_parts.append(f"### Notes\n\n{notes_text}")

        slide_content = "\n\n".join(part.strip() for part in slide_parts if part and part.strip()).strip()
        if slide_content:
            content += slide_content

        if i < len(presentation.slides):
            content += "---\n\n"

    return content.strip(), extracted_images
