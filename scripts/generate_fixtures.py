"""Regenerate the small, synthetic conversion corpus; run from the repository root."""
import io
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
out = root / 'tests' / 'fixtures'
out.mkdir(exist_ok=True)
shutil.copyfile(root / 'dsh-plugin' / 'tests' / 'fixtures' / 'sample.docx', out / 'report.docx')
book = Workbook()
sheet = book.active
sheet.title = 'Budget'
sheet.append(['Item', 'Cost', 'Cached total'])
sheet.append(['First', 10])
sheet.append(['Second', 20])
sheet.append(['Total', '=SUM(B2:B3)', '=SUM(B2:B3)'])
buffer = io.BytesIO()
book.save(buffer)
with ZipFile(buffer) as source, ZipFile(out / 'budget.xlsx', 'w', ZIP_DEFLATED) as target:
    for name in source.namelist():
        data = source.read(name)
        if name == 'xl/worksheets/sheet1.xml':
            ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
            xml = ET.fromstring(data)
            cell = xml.find(f'.//{{{ns}}}c[@r="C4"]')
            if cell is None or cell.find(f'{{{ns}}}f') is None:
                raise ValueError('Fixture requires a formula in C4')
            value = cell.find(f'{{{ns}}}v')
            if value is None:
                value = ET.SubElement(cell, f'{{{ns}}}v')
            value.text = '30'
            data = ET.tostring(xml, encoding='utf-8')
        target.writestr(name, data)
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = 'Launch checklist'
slide.placeholders[1].text = 'Verify source\nCheck output\nPublish only after review'
slides.save(out / 'slides.pptx')

stream = b'BT /F1 12 Tf 72 720 Td (PDF fixture text) Tj ET'
objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
           b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
           b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
           b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
           b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
pdf = bytearray(b'%PDF-1.4\n')
offsets = [0]
for index, body in enumerate(objects, 1):
    offsets.append(len(pdf))
    pdf.extend(f'{index} 0 obj\n'.encode() + body + b'\nendobj\n')
xref = len(pdf)
pdf.extend(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode())
for offset in offsets[1:]:
    pdf.extend(f'{offset:010d} 00000 n \n'.encode())
pdf.extend(f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
(out / 'report.pdf').write_bytes(pdf)
